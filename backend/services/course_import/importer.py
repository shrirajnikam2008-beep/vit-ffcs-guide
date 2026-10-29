"""
Excel & CSV Course Allocation Importer using openpyxl and csv (zero C-extension DLL issues).
Supports:
- Automatic & manual column mapping
- Multi-worksheet reading (Course Summary, Original Course List, Combined Faculty Options, Per-Course sheets, Unnamed C1/C2 Applied Chemistry sheet)
- Duplicate & cross-course anomaly detection
- Missing field & unresolved slot flagging
- Source worksheet and row traceability
"""
import csv
import io
import json
import os
import re
from typing import Dict, List, Optional, Tuple, Any
from openpyxl import load_workbook
from sqlalchemy.orm import Session
from backend.models.database import (
    CourseModel,
    CourseSectionModel,
    ImportBatchLogModel,
    BASE_DIR,
)
from backend.services.no8am_scraper.service import normalize_faculty_name
from backend.services.slot_parser.engine import parse_slot_string, load_active_slot_map

KNOWN_VIT_COURSE_METADATA: Dict[str, Dict[str, Any]] = {
    "applied chemistry": {
        "course_code": "BACHY105",
        "course_title": "Applied Chemistry",
        "curriculum_category": "UCC",
        "credits": 4.0,
        "course_type": "ETL",
        "lecture_hours": 3,
        "tutorial_hours": 0,
        "practical_hours": 2,
        "project_hours": 0,
        "employee_school": "SAS",
        "verification_status": "VERIFIED",
    },
    "structured and object-oriented programming": {
        "course_code": "BACSE104",
        "course_title": "Structured and Object-Oriented Programming",
        "curriculum_category": "PCC",
        "credits": 4.0,
        "course_type": "ETL",
        "lecture_hours": 2,
        "tutorial_hours": 0,
        "practical_hours": 4,
        "project_hours": 0,
        "employee_school": "SCORE",
        "verification_status": "VERIFIED",
    },
    "structured and object oriented programming": {
        "course_code": "BACSE104",
        "course_title": "Structured and Object-Oriented Programming",
        "curriculum_category": "PCC",
        "credits": 4.0,
        "course_type": "ETL",
        "lecture_hours": 2,
        "tutorial_hours": 0,
        "practical_hours": 4,
        "project_hours": 0,
        "employee_school": "SCORE",
        "verification_status": "VERIFIED",
    },
    "technical english communication": {
        "course_code": "BAENG101",
        "course_title": "Technical English Communication",
        "curriculum_category": "UCC",
        "credits": 4.0,
        "course_type": "ETL",
        "lecture_hours": 3,
        "tutorial_hours": 0,
        "practical_hours": 2,
        "project_hours": 0,
        "employee_school": "SSL",
        "verification_status": "MISSING_DATA",
    },
    "india studies": {
        "course_code": "BAHUM101",
        "course_title": "India Studies",
        "curriculum_category": "UCC",
        "credits": 1.0,
        "course_type": "OC",
        "lecture_hours": 0,
        "tutorial_hours": 0,
        "practical_hours": 0,
        "project_hours": 0,
        "employee_school": "SSL",
        "verification_status": "VERIFIED",
    },
    "digital logic design and computer organization": {
        "course_code": "BAITE101",
        "course_title": "Digital Logic Design and Computer Organization",
        "curriculum_category": "PCC",
        "credits": 4.0,
        "course_type": "ETL",
        "lecture_hours": 3,
        "tutorial_hours": 0,
        "practical_hours": 2,
        "project_hours": 0,
        "employee_school": "SCORE",
        "verification_status": "VERIFIED",
    },
    "data structures and algorithms": {
        "course_code": "BAITE102",
        "course_title": "Data Structures and Algorithms",
        "curriculum_category": "PCC",
        "credits": 4.0,
        "course_type": "ETL",
        "lecture_hours": 3,
        "tutorial_hours": 0,
        "practical_hours": 2,
        "project_hours": 0,
        "employee_school": "SCORE",
        "verification_status": "VERIFIED",
    },
    "principles of software engineering": {
        "course_code": "BAITE103",
        "course_title": "Principles of Software Engineering",
        "curriculum_category": "PCC",
        "credits": 4.0,
        "course_type": "ETL",
        "lecture_hours": 3,
        "tutorial_hours": 0,
        "practical_hours": 2,
        "project_hours": 0,
        "employee_school": "SCORE",
        "verification_status": "VERIFIED",
    },
    "probability and statistics": {
        "course_code": "BAMAT207",
        "course_title": "Probability and Statistics",
        "curriculum_category": "PCC",
        "credits": 4.0,
        "course_type": "ETL",
        "lecture_hours": 3,
        "tutorial_hours": 0,
        "practical_hours": 2,
        "project_hours": 0,
        "employee_school": "SAS",
        "verification_status": "VERIFIED",
    },
    "german level i": {
        "course_code": "BAGER101",
        "course_title": "German Level I",
        "curriculum_category": "University Elective",
        "credits": 2.0,
        "course_type": "ETL",
        "lecture_hours": 1,
        "tutorial_hours": 0,
        "practical_hours": 2,
        "project_hours": 0,
        "employee_school": "SSL",
        "verification_status": "INFERRED",
    },
    "database information systems": {
        "course_code": "BAITE202",
        "course_title": "Database Information Systems",
        "curriculum_category": "Programme Elective",
        "credits": 4.0,
        "course_type": "ETL",
        "lecture_hours": 3,
        "tutorial_hours": 0,
        "practical_hours": 2,
        "project_hours": 0,
        "employee_school": "SCORE",
        "verification_status": "INFERRED",
    },
}

CHEMISTRY_LAB_VENUES = {"PRP408A", "PRPG07"}


def norm_title(title: str) -> str:
    if not title:
        return ""
    return " ".join(title.strip().lower().split())


def read_file_to_sheets(file_bytes: bytes, filename: str) -> Dict[str, Tuple[List[str], List[Dict[str, str]]]]:
    """
    Returns dict mapping sheet_name -> (columns_list, list_of_row_dicts).
    """
    sheets: Dict[str, Tuple[List[str], List[Dict[str, str]]]] = {}
    if filename.lower().endswith(".csv"):
        text = file_bytes.decode("utf-8", errors="replace")
        reader = list(csv.reader(io.StringIO(text)))
        if not reader:
            return {"CSV Import": ([], [])}
        headers = [str(h).strip() for h in reader[0]]
        rows = []
        for r in reader[1:]:
            row_dict = {
                headers[i]: (str(r[i]) if i < len(r) and r[i] is not None else "")
                for i in range(len(headers))
            }
            rows.append(row_dict)
        sheets["CSV Import"] = (headers, rows)
    else:
        wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
        for s_name in wb.sheetnames:
            ws = wb[s_name]
            all_rows = list(ws.iter_rows(values_only=True))
            if not all_rows:
                sheets[s_name] = ([], [])
                continue
            headers = [
                str(c).strip() if c is not None else f"Col_{idx}"
                for idx, c in enumerate(all_rows[0])
            ]
            rows = []
            for r in all_rows[1:]:
                row_dict = {
                    headers[i]: (str(r[i]) if i < len(r) and r[i] is not None else "")
                    for i in range(len(headers))
                }
                rows.append(row_dict)
            sheets[s_name] = (headers, rows)
    return sheets


def detect_sheet_kind_and_columns(headers: List[str], sheet_name: str) -> Dict[str, Any]:
    cols_lower = {str(c).strip().lower(): str(c) for c in headers}

    if "note" in cols_lower and len(cols_lower) <= 2:
        return {"kind": "notes", "mapping": {"note": cols_lower["note"]}}

    if ("course code" in cols_lower or "course_code" in cols_lower) and (
        "credits" in cols_lower or "lecture slot" in cols_lower or "slot" in cols_lower
    ):
        mapping = {
            "course_code": cols_lower.get("course code") or cols_lower.get("course_code") or "",
            "course_title": cols_lower.get("course title") or cols_lower.get("course_title") or "",
            "curriculum_category": cols_lower.get("curriculum category") or cols_lower.get("type") or "",
            "batch": cols_lower.get("batch") or "",
            "course_type": cols_lower.get("course type") or cols_lower.get("course_type") or "",
            "credits": cols_lower.get("credits") or "",
            "lecture_hours": cols_lower.get("lecture hours") or "",
            "tutorial_hours": cols_lower.get("tutorial hours") or "",
            "practical_hours": cols_lower.get("practical hours") or "",
            "project_hours": cols_lower.get("project hours") or "",
            "employee_school": cols_lower.get("employee school") or "",
            "strength": cols_lower.get("strength") or "",
            "lecture_slot": cols_lower.get("lecture slot") or cols_lower.get("slot") or "",
            "lab_slot": cols_lower.get("lab/tutorial slot") or cols_lower.get("slot1") or "",
        }
        return {"kind": "course_summary", "mapping": mapping}

    if ("slot detail" in cols_lower or "slot" in cols_lower) and ("faculty" in cols_lower or "venue" in cols_lower):
        mapping = {
            "course": cols_lower.get("course") or cols_lower.get("course title") or cols_lower.get("course code") or "",
            "slot_detail": cols_lower.get("slot detail") or cols_lower.get("slot") or "",
            "venue": cols_lower.get("venue") or cols_lower.get("room") or "",
            "faculty": cols_lower.get("faculty") or cols_lower.get("faculty name") or cols_lower.get("professor") or "",
            "course_type": cols_lower.get("course type") or cols_lower.get("type") or "",
        }
        return {"kind": "faculty_offerings", "mapping": mapping}

    return {"kind": "unknown", "mapping": {str(c): str(c) for c in headers}}


def map_course_type_to_class_type(course_type_code: str, slot_str: str) -> str:
    ct = (course_type_code or "").strip().upper()
    s = (slot_str or "").strip().upper()
    if ct == "OC" or s in ("OC", "NIL"):
        return "online"
    if ct == "ELA" or s.startswith("L"):
        return "lab"
    if ct == "ETH":
        return "theory"
    if ct == "ETL":
        return "theory_and_lab"
    return "other"


def preview_workbook_bytes(
    file_bytes: bytes,
    filename: str,
    custom_mappings: Optional[Dict[str, Dict[str, str]]] = None,
) -> Dict[str, Any]:
    sheets_dict = read_file_to_sheets(file_bytes, filename)
    slot_map = load_active_slot_map(None, "vellore")
    sheet_summaries = []
    issues = []
    total_rows = 0

    for s_name, (headers, rows) in sheets_dict.items():
        detected = detect_sheet_kind_and_columns(headers, s_name)
        if custom_mappings and s_name in custom_mappings:
            detected["mapping"].update(custom_mappings[s_name])

        total_rows += len(rows)
        sheet_summaries.append(
            {
                "sheet_name": s_name,
                "kind": detected["kind"],
                "row_count": len(rows),
                "columns": headers,
                "detected_mapping": detected["mapping"],
                "sample_rows": rows[:15],
            }
        )

        if detected["kind"] == "faculty_offerings":
            m = detected["mapping"]
            seen_keys = set()
            for idx, row in enumerate(rows):
                row_num = idx + 2
                c_raw = str(row.get(m.get("course", ""), "")).strip()
                s_raw = str(row.get(m.get("slot_detail", ""), "")).strip()
                v_raw = str(row.get(m.get("venue", ""), "")).strip()
                f_raw = str(row.get(m.get("faculty", ""), "")).strip()
                t_raw = str(row.get(m.get("course_type", ""), "")).strip()

                if not c_raw and ("C1" in s_raw or "C2" in s_raw or v_raw in CHEMISTRY_LAB_VENUES):
                    c_raw = "Applied Chemistry"
                    issues.append(
                        {
                            "severity": "INFO",
                            "type": "UNNAMED_COURSE_ASSIGNED",
                            "sheet": s_name,
                            "row": row_num,
                            "message": "Unnamed course row assigned to 'Applied Chemistry' per user specification.",
                        }
                    )

                if not f_raw:
                    issues.append(
                        {
                            "severity": "WARNING",
                            "type": "MISSING_FACULTY",
                            "sheet": s_name,
                            "row": row_num,
                            "message": f"Missing faculty name for {c_raw or 'Unknown Course'} ({s_raw}).",
                        }
                    )
                if not v_raw or (v_raw.upper() == "NIL" and t_raw.upper() != "OC"):
                    issues.append(
                        {
                            "severity": "WARNING",
                            "type": "MISSING_VENUE",
                            "sheet": s_name,
                            "row": row_num,
                            "message": f"Missing venue for {c_raw} ({f_raw}, {s_raw}).",
                        }
                    )

                key = (norm_title(c_raw), s_raw.upper(), v_raw.upper(), normalize_faculty_name(f_raw), t_raw.upper())
                if key in seen_keys:
                    issues.append(
                        {
                            "severity": "WARNING",
                            "type": "DUPLICATE_ROW",
                            "sheet": s_name,
                            "row": row_num,
                            "message": f"Duplicate offering in sheet '{s_name}': {c_raw} | {f_raw} | {s_raw} | {v_raw}.",
                        }
                    )
                seen_keys.add(key)

                if norm_title(c_raw) == "probability and statistics" and v_raw.upper() in CHEMISTRY_LAB_VENUES:
                    issues.append(
                        {
                            "severity": "WARNING",
                            "type": "CROSS_COURSE_ANOMALY",
                            "sheet": s_name,
                            "row": row_num,
                            "message": (
                                f"Row in '{s_name}' lists Chemistry lab venue {v_raw} ({f_raw}, {s_raw}) under "
                                "'Probability and Statistics' (matches Applied Chemistry ELA block)."
                            ),
                        }
                    )

        elif detected["kind"] == "course_summary":
            m = detected["mapping"]
            for idx, row in enumerate(rows):
                row_num = idx + 2
                code = str(row.get(m.get("course_code", ""), "")).strip()
                l_slot = str(row.get(m.get("lecture_slot", ""), "")).strip()
                lab_slot = str(row.get(m.get("lab_slot", ""), "")).strip()
                for s_val in (l_slot, lab_slot):
                    if s_val:
                        parsed = parse_slot_string(s_val, slot_map)
                        if parsed.status == "UNRESOLVED":
                            issues.append(
                                {
                                    "severity": "WARNING",
                                    "type": "UNRESOLVED_BATCH_SLOT",
                                    "sheet": s_name,
                                    "row": row_num,
                                    "message": f"Course {code} has unresolved batch slot '{s_val}': {parsed.explanation}",
                                }
                            )

    return {
        "filename": filename,
        "total_sheets": len(sheets_dict),
        "total_rows": total_rows,
        "sheets": sheet_summaries,
        "issues_count": len(issues),
        "issues": issues,
    }


def import_workbook_to_db(
    db: Session,
    file_bytes: bytes,
    filename: str,
    custom_mappings: Optional[Dict[str, Dict[str, str]]] = None,
    clear_existing: bool = True,
) -> Dict[str, Any]:
    preview = preview_workbook_bytes(file_bytes, filename, custom_mappings)
    sheets_dict = read_file_to_sheets(file_bytes, filename)

    if clear_existing:
        db.query(CourseSectionModel).delete()
        db.query(CourseModel).delete()
        db.commit()

    slot_map = load_active_slot_map(db, "vellore")

    courses_by_code: Dict[str, Dict[str, Any]] = {}
    title_to_code: Dict[str, str] = {}
    for norm_t, meta in KNOWN_VIT_COURSE_METADATA.items():
        title_to_code[norm_t] = meta["course_code"]

    for s_name, (headers, rows) in sheets_dict.items():
        detected = detect_sheet_kind_and_columns(headers, s_name)
        if custom_mappings and s_name in custom_mappings:
            detected["mapping"].update(custom_mappings[s_name])
        if detected["kind"] != "course_summary":
            continue

        m = detected["mapping"]
        for idx, row in enumerate(rows):
            row_num = idx + 2
            code = str(row.get(m.get("course_code", ""), "")).strip()
            title = str(row.get(m.get("course_title", ""), "")).strip()
            if not code or code.upper() == "COURSE CODE":
                continue

            cat = str(row.get(m.get("curriculum_category", ""), "")).strip() or "PCC"
            batch = str(row.get(m.get("batch", ""), "")).strip()
            ctype = str(row.get(m.get("course_type", ""), "")).strip() or "ETL"
            try:
                credits = float(row.get(m.get("credits", ""), 4) or 4)
            except ValueError:
                credits = 4.0
            try:
                lh = int(float(row.get(m.get("lecture_hours", ""), 0) or 0))
                th = int(float(row.get(m.get("tutorial_hours", ""), 0) or 0))
                ph = int(float(row.get(m.get("practical_hours", ""), 0) or 0))
                jh = int(float(row.get(m.get("project_hours", ""), 0) or 0))
            except ValueError:
                lh, th, ph, jh = 3, 0, 2, 0

            school = str(row.get(m.get("employee_school", ""), "")).strip() or None
            strength_raw = str(row.get(m.get("strength", ""), "")).strip()
            strength = int(float(strength_raw)) if strength_raw and strength_raw.replace(".", "").isdigit() else None
            l_slot = str(row.get(m.get("lecture_slot", ""), "")).strip()
            lab_slot = str(row.get(m.get("lab_slot", ""), "")).strip()

            norm_t = norm_title(title)
            if norm_t:
                title_to_code[norm_t] = code

            if code not in courses_by_code:
                known = KNOWN_VIT_COURSE_METADATA.get(norm_t, {})
                courses_by_code[code] = {
                    "course_code": code,
                    "course_title": title or known.get("course_title", code),
                    "semester": "Winter Semester 2026-27",
                    "campus": "VIT Vellore",
                    "curriculum_category": cat,
                    "credits": credits,
                    "course_type": ctype,
                    "lecture_hours": lh,
                    "tutorial_hours": th,
                    "practical_hours": ph,
                    "project_hours": jh,
                    "employee_school": school or known.get("employee_school"),
                    "theory_slots": [],
                    "lab_slots": [],
                    "batches": [],
                    "verification_status": "VERIFIED",
                    "verification_note": None,
                    "source_sheet": s_name,
                    "source_row": row_num,
                }
            else:
                if school and not courses_by_code[code]["employee_school"]:
                    courses_by_code[code]["employee_school"] = school

            entry = courses_by_code[code]
            if l_slot and l_slot not in entry["theory_slots"]:
                entry["theory_slots"].append(l_slot)
            if lab_slot and lab_slot not in entry["lab_slots"]:
                entry["lab_slots"].append(lab_slot)

            if batch and not any(b["batch"] == batch for b in entry["batches"]):
                entry["batches"].append(
                    {
                        "batch": batch,
                        "course_type": ctype,
                        "lecture_slot": l_slot,
                        "lab_slot": lab_slot,
                        "strength": strength,
                    }
                )

    sections_to_insert: List[CourseSectionModel] = []
    seen_primary_offerings: Dict[Tuple[str, str, str, str, str], int] = {}
    duplicates_count = 0
    anomalies_count = 0

    for s_name, (headers, rows) in sheets_dict.items():
        detected = detect_sheet_kind_and_columns(headers, s_name)
        if custom_mappings and s_name in custom_mappings:
            detected["mapping"].update(custom_mappings[s_name])
        if detected["kind"] != "faculty_offerings":
            continue

        m = detected["mapping"]
        for idx, row in enumerate(rows):
            row_num = idx + 2
            c_raw = str(row.get(m.get("course", ""), "")).strip()
            s_raw = str(row.get(m.get("slot_detail", ""), "")).strip()
            v_raw = str(row.get(m.get("venue", ""), "")).strip()
            f_raw = str(row.get(m.get("faculty", ""), "")).strip()
            t_raw = str(row.get(m.get("course_type", ""), "")).strip()

            if not c_raw and not s_raw and not f_raw:
                continue
            if c_raw.lower() == "course" and s_raw.lower() == "slot detail":
                continue

            if not c_raw:
                c_raw = "Applied Chemistry"

            norm_c = norm_title(c_raw)
            code = title_to_code.get(norm_c)
            if not code:
                if c_raw.upper() in courses_by_code:
                    code = c_raw.upper()
                else:
                    code = f"VIT_{re.sub(r'[^A-Z0-9]', '', c_raw.upper())[:8]}"
                    title_to_code[norm_c] = code

            if code not in courses_by_code:
                known = KNOWN_VIT_COURSE_METADATA.get(norm_c, {})
                courses_by_code[code] = {
                    "course_code": code,
                    "course_title": known.get("course_title", c_raw),
                    "semester": "Winter Semester 2026-27",
                    "campus": "VIT Vellore",
                    "curriculum_category": known.get("curriculum_category", "University Elective"),
                    "credits": float(known.get("credits", 3.0)),
                    "course_type": known.get("course_type", "ETL"),
                    "lecture_hours": int(known.get("lecture_hours", 2)),
                    "tutorial_hours": int(known.get("tutorial_hours", 0)),
                    "practical_hours": int(known.get("practical_hours", 2)),
                    "project_hours": int(known.get("project_hours", 0)),
                    "employee_school": known.get("employee_school", "VIT"),
                    "theory_slots": [],
                    "lab_slots": [],
                    "batches": [],
                    "verification_status": known.get("verification_status", "INFERRED"),
                    "verification_note": (
                        "Course metadata inferred from VIT Winter 2026-27 catalog because this course "
                        "appeared in the Faculty Options sheet without a row in the Batch Course Summary."
                    ),
                    "source_sheet": s_name,
                    "source_row": row_num,
                }

            c_entry = courses_by_code[code]
            norm_fac = normalize_faculty_name(f_raw)
            class_type = map_course_type_to_class_type(t_raw, s_raw)

            if class_type == "lab" and s_raw and s_raw not in c_entry["lab_slots"]:
                c_entry["lab_slots"].append(s_raw)
            elif class_type in ("theory", "online") and s_raw and s_raw not in c_entry["theory_slots"]:
                c_entry["theory_slots"].append(s_raw)

            dedup_key = (code, s_raw.upper(), v_raw.upper(), norm_fac, t_raw.upper())
            is_dup = dedup_key in seen_primary_offerings
            if is_dup:
                duplicates_count += 1
            else:
                seen_primary_offerings[dedup_key] = row_num

            is_anom = False
            anom_note = None
            if code == "BAMAT207" and v_raw.upper() in CHEMISTRY_LAB_VENUES:
                is_anom = True
                anomalies_count += 1
                anom_note = (
                    f"Chemistry laboratory venue ({v_raw}) and SAS-Chemistry faculty ({f_raw}) "
                    "appeared under Probability and Statistics in the Combined sheet, matching the "
                    "Applied Chemistry (BACHY105) lab table. Preserved for audit; hidden from active "
                    "BAMAT207 selection by default."
                )

            parsed_slot = parse_slot_string(s_raw, slot_map)
            ver_status = "VERIFIED"
            if parsed_slot.status == "UNRESOLVED":
                ver_status = "UNRESOLVED"
            elif is_anom:
                ver_status = "INFERRED"

            sections_to_insert.append(
                CourseSectionModel(
                    course_code=code,
                    course_title=c_entry["course_title"],
                    faculty_name=f_raw,
                    normalized_faculty_name=norm_fac,
                    slot_combination=s_raw,
                    venue=v_raw,
                    class_type=class_type,
                    course_type_code=t_raw or "ETH",
                    batch=None,
                    semester="Winter Semester 2026-27",
                    campus="VIT Vellore",
                    source_sheet=s_name,
                    source_row=row_num,
                    is_duplicate=is_dup,
                    is_anomaly=is_anom,
                    anomaly_note=anom_note,
                    is_active=not (is_dup or is_anom),
                    verification_status=ver_status,
                )
            )

    active_counts_by_code: Dict[str, int] = {}
    for sec in sections_to_insert:
        if sec.is_active:
            active_counts_by_code[sec.course_code] = active_counts_by_code.get(sec.course_code, 0) + 1

    for code, c_entry in courses_by_code.items():
        if active_counts_by_code.get(code, 0) == 0:
            c_entry["verification_status"] = "MISSING_DATA"
            c_entry["verification_note"] = (
                f"{c_entry['course_title']} ({code}) is present in the Batch Course Summary "
                f"(slots: {', '.join(c_entry['theory_slots'] + c_entry['lab_slots'])}) "
                "but has no individual faculty allocation rows in the Faculty Options sheets. "
                "Placeholder batch sections have been created so you can still plan around it."
            )
            for b in c_entry["batches"]:
                if b["lecture_slot"]:
                    sections_to_insert.append(
                        CourseSectionModel(
                            course_code=code,
                            course_title=c_entry["course_title"],
                            faculty_name=f"TBA ({b['batch']} Theory)",
                            normalized_faculty_name=f"tba {b['batch'].lower()} theory",
                            slot_combination=b["lecture_slot"],
                            venue="TBA",
                            class_type="theory",
                            course_type_code="ETH",
                            batch=b["batch"],
                            semester="Winter Semester 2026-27",
                            campus="VIT Vellore",
                            source_sheet=c_entry["source_sheet"] or "Course Summary",
                            source_row=c_entry["source_row"] or 1,
                            is_duplicate=False,
                            is_anomaly=False,
                            anomaly_note="Placeholder created from Batch Course Summary (no faculty list provided).",
                            is_active=True,
                            verification_status="UNRESOLVED",
                        )
                    )
                if b["lab_slot"]:
                    sections_to_insert.append(
                        CourseSectionModel(
                            course_code=code,
                            course_title=c_entry["course_title"],
                            faculty_name=f"TBA ({b['batch']} Lab)",
                            normalized_faculty_name=f"tba {b['batch'].lower()} lab",
                            slot_combination=b["lab_slot"],
                            venue="TBA",
                            class_type="lab",
                            course_type_code="ELA",
                            batch=b["batch"],
                            semester="Winter Semester 2026-27",
                            campus="VIT Vellore",
                            source_sheet=c_entry["source_sheet"] or "Course Summary",
                            source_row=c_entry["source_row"] or 1,
                            is_duplicate=False,
                            is_anomaly=False,
                            anomaly_note="Placeholder created from Batch Course Summary (no faculty list provided).",
                            is_active=True,
                            verification_status="UNRESOLVED",
                        )
                    )

    for code, c_entry in sorted(courses_by_code.items()):
        db.add(
            CourseModel(
                course_code=c_entry["course_code"],
                course_title=c_entry["course_title"],
                semester=c_entry["semester"],
                campus=c_entry["campus"],
                curriculum_category=c_entry["curriculum_category"],
                credits=c_entry["credits"],
                course_type=c_entry["course_type"],
                lecture_hours=c_entry["lecture_hours"],
                tutorial_hours=c_entry["tutorial_hours"],
                practical_hours=c_entry["practical_hours"],
                project_hours=c_entry["project_hours"],
                employee_school=c_entry["employee_school"],
                theory_slots_json=json.dumps(c_entry["theory_slots"]),
                lab_slots_json=json.dumps(c_entry["lab_slots"]),
                batches_json=json.dumps(c_entry["batches"]),
                verification_status=c_entry["verification_status"],
                verification_note=c_entry["verification_note"],
                source_sheet=c_entry["source_sheet"],
                source_row=c_entry["source_row"],
            )
        )

    for sec in sections_to_insert:
        db.add(sec)

    log_entry = ImportBatchLogModel(
        filename=filename,
        sheets_detected_json=json.dumps([s["sheet_name"] for s in preview["sheets"]]),
        total_rows_read=preview["total_rows"],
        courses_imported=len(courses_by_code),
        sections_imported=len(sections_to_insert),
        duplicates_detected=duplicates_count,
        anomalies_detected=anomalies_count,
        missing_fields_count=sum(
            1 for i in preview["issues"] if i["type"] in ("MISSING_FACULTY", "MISSING_VENUE", "UNRESOLVED_BATCH_SLOT")
        ),
        issues_json=json.dumps(preview["issues"]),
        raw_preview_json=json.dumps(preview["sheets"]),
    )
    db.add(log_entry)
    db.commit()

    return {
        "status": "success",
        "filename": filename,
        "courses_imported": len(courses_by_code),
        "sections_imported": len(sections_to_insert),
        "active_sections": sum(1 for s in sections_to_insert if s.is_active),
        "duplicates_detected": duplicates_count,
        "anomalies_detected": anomalies_count,
        "issues": preview["issues"],
    }


def seed_default_workbook_if_empty(db: Session) -> None:
    if db.query(CourseModel).count() > 0:
        return
    wb_path = os.path.join(BASE_DIR, "data", "vit_winter_2026_27_allocation.xlsx")
    if not os.path.exists(wb_path):
        from backend.data.generate_sample_workbook import build_workbook
        build_workbook(wb_path)
    with open(wb_path, "rb") as f:
        data = f.read()
    import_workbook_to_db(db, data, "vit_winter_2026_27_allocation.xlsx", clear_existing=True)
